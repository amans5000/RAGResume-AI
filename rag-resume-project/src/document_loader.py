import os
import fitz
from docx import Document


def load_pdf(file_path):
    """
    Extract text from a PDF file.
    """

    doc = fitz.open(file_path)

    text = ""

    for page in doc:
        page_text = page.get_text()

        if page_text.strip():
            text += page_text
            text += "\n"

    doc.close()

    return text


def load_docx(file_path):
    """
    Extract text from a DOCX file.
    """

    doc = Document(file_path)

    text = ""

    # Extract paragraphs
    for paragraph in doc.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text
            text += "\n"

    # Extract tables
    for table in doc.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:
                row_text.append(cell.text.strip())

            text += " | ".join(row_text)
            text += "\n"

    return text


def load_document(file_path):
    """
    Automatically detect PDF or DOCX
    and extract its text.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        return load_pdf(file_path)

    elif extension == ".docx":

        return load_docx(file_path)

    else:

        raise ValueError(
            "Unsupported file format. "
            "Only PDF and DOCX are supported."
        )